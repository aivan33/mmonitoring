"""Loader for the FaradaIC MR (Management Reporting) workbook.

The MR is the master DATEV-derived workbook with consolidated P&L, CF, and
BS sheets (Germany + Serbia, FX-converted upstream). Its layout differs
from the canonical taxonomi format — header in row 2, monthly columns
spanning years — so it has its own loader. Output feeds
``core/mr_to_taxonomi.py`` which writes a new month's column into a copy
of the prior taxonomi-actual xlsx.

See ``clients/farada/MR_LAYOUT.md`` for sheet coordinates and
``clients/farada/mapping.yaml`` for the row-to-taxonomi mapping.
"""

from __future__ import annotations

import datetime as dt
import logging
from pathlib import Path

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


# (sheet name, label column 1-indexed) per statement code.
_SHEET_META: dict[str, tuple[str, int]] = {
    "IS": ("P&L", 1),
    "CF": ("CF", 2),
    "BS": ("BS", 2),
}

# Header row is fixed at row 2 across all consumed MR sheets.
_HEADER_ROW = 2

_MAPPING_KEY: dict[str, str] = {
    "IS": "mapping_is",
    "CF": "mapping_cf",
    "BS": "mapping_bs",
}

# Cell strings treated as "no value".
_BLANK_STRINGS = frozenset({"", "-", "—", "n/a", "N/A"})


def extract_month(
    mr_path: str | Path,
    mapping: dict,
    year: int,
    month: int,
    statement: str,
) -> dict[tuple[str, str, str], float | None]:
    """Extract one month's values from an MR sheet, mapped onto taxonomi keys.

    Args:
        mr_path: path to the MR workbook.
        mapping: parsed ``mapping.yaml`` dict; must contain the per-statement
            list keyed by ``mapping_is`` / ``mapping_cf`` / ``mapping_bs``.
        year, month: the target period.
        statement: ``'IS' | 'CF' | 'BS'``.

    Returns:
        ``dict`` mapping ``(data, grp, subgroup)`` → ``float`` or ``None``.
        ``None`` is emitted when the cell is empty / a dash placeholder, or
        when the mapping entry has ``mr_row=null`` (derived/non-MR rows).

    Behavior:
        - Resolves the target column by scanning ``row 2`` for a date cell
          matching ``(year, month)``. Raises ``ValueError`` if absent.
        - For each entry: reads the configured row's label; if it matches
          ``mr_label``, uses the configured row. Otherwise searches the
          label column for ``mr_label`` and uses that row instead, emitting
          a ``logging.warning`` with both row indices. If the label can't
          be found anywhere, the entry's value is ``None`` and a warning
          is emitted.
    """
    if statement not in _SHEET_META:
        raise ValueError(
            f"statement={statement!r}; expected one of {list(_SHEET_META)}"
        )

    sheet_name, label_col = _SHEET_META[statement]
    entries = mapping[_MAPPING_KEY[statement]]

    wb = load_workbook(mr_path, data_only=True)
    try:
        ws = wb[sheet_name]
        target_col = _find_month_column(ws, year, month, statement)
        return _extract(ws, entries, label_col, target_col, statement)
    finally:
        wb.close()


def _find_month_column(ws, year: int, month: int, statement: str) -> int:
    for c in range(1, ws.max_column + 1):
        v = ws.cell(_HEADER_ROW, c).value
        if isinstance(v, dt.date) and v.year == year and v.month == month:
            return c
    raise ValueError(
        f"{statement} sheet: no header column for {year}-{month:02d}"
    )


def _extract(ws, entries, label_col: int, target_col: int, statement: str):
    out: dict[tuple[str, str, str], float | None] = {}
    for entry in entries:
        key = (entry["data"], entry["grp"], entry["subgroup"])
        mr_row = entry["mr_row"]
        if mr_row is None:
            out[key] = None
            continue
        row = _resolve_row(
            ws, mr_row, entry["mr_label"], label_col, statement,
        )
        if row is None:
            out[key] = None
            continue
        value = _coerce_value(ws.cell(row, target_col).value)
        if value is not None:
            sign = entry.get("sign", 1)
            value = value * sign
        out[key] = value
    return out


def _resolve_row(
    ws, mr_row: int, mr_label: str, label_col: int, statement: str,
) -> int | None:
    actual = ws.cell(mr_row, label_col).value
    if actual == mr_label:
        return mr_row
    for r in range(_HEADER_ROW + 1, ws.max_row + 1):
        if ws.cell(r, label_col).value == mr_label:
            logger.warning(
                "%s: configured row %d has label %r, but %r found at row %d "
                "— using row %d. Update mapping.yaml.",
                statement, mr_row, actual, mr_label, r, r,
            )
            return r
    logger.warning(
        "%s: configured row %d has label %r, expected %r — label not found "
        "elsewhere in sheet. Leaving cell null.",
        statement, mr_row, actual, mr_label,
    )
    return None


def _coerce_value(raw):
    if raw is None:
        return None
    if isinstance(raw, str):
        if raw.strip() in _BLANK_STRINGS:
            return None
        try:
            return float(raw)
        except ValueError:
            return None
    return float(raw)
